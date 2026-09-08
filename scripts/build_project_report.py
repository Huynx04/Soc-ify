#!/usr/bin/env python3
"""Generate a polished Word (.docx) project report for SOC-ify."""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

OUT = os.path.join(os.path.dirname(__file__), "..", "reports", "SOC-ify_Project_Report.docx")
DARK = RGBColor(0x1F, 0x3A, 0x5F)
ACCENT = RGBColor(0xC2, 0x3B, 0x22)  # SOC red accent

doc = Document()

# ---- Base style ----
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)

def h1(text):
    p = doc.add_heading(text, level=1)
    for r in p.runs:
        r.font.color.rgb = DARK
    return p

def h2(text):
    p = doc.add_heading(text, level=2)
    for r in p.runs:
        r.font.color.rgb = DARK
    return p

def para(text, bold=False, color=None, size=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    if color: r.font.color.rgb = color
    if size: r.font.size = Pt(size)
    return p

def bullet(text):
    doc.add_paragraph(text, style="List Bullet")

# ---- Cover ----
t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("SOC-ify")
r.font.size = Pt(40); r.bold = True; r.font.color.rgb = DARK
sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("Mini Security Operations Center trên Azure + Elasticsearch")
r.font.size = Pt(16); r.font.color.rgb = ACCENT
meta = doc.add_paragraph(); meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = meta.add_run("Báo cáo tổng thể dự án  •  Elasticsearch 8.12 + Kibana  •  MIT License")
r.font.size = Pt(10)
doc.add_paragraph()

# ---- 1. Tổng quan ----
h1("1. Tổng quan dự án")
para("SOC-ify là một Security Operations Center (SOC) thu nhỏ, có thể tái tạo "
     "(reproducible), xây dựng trên một Azure VM phơi bày có chủ đích, với mục tiêu "
     "biến log truy cập nginx thành một pipeline giám sát bảo mật hoàn chỉnh: "
     "Phát hiện → Triage → Báo cáo (detect → triage → report), kèm dashboard Kibana trực tiếp.")
para("Điểm đặc biệt: hệ thống hoàn toàn tương thích giấy phép Basic của Elastic "
     "(không phụ thuộc các tính năng trả phí như Elastic Security / Watcher).", bold=True)

# ---- 2. Kiến trúc ----
h1("2. Kiến trúc hệ thống")
para("Luồng dữ liệu: Internet → Azure VM → nginx (+ ModSecurity CRS + geo-block VN-only). "
     "Filebeat thu thập 3 nguồn log rồi đẩy về Elasticsearch:")
bullet("nginx access.log  → index nginx-access-*")
bullet("UFW firewall log  → index fire-ufw-*")
bullet("SSH auth log      → index auth-*")
para("Đường ống xử lý:")
bullet("Ingest pipeline (nginx-soc-enrich.json) tự phân loại mỗi request thành class tấn công "
       "(SQLi, XSS, path traversal, command injection, SSRF, scanner...) kèm severity + confidence.")
bullet("scripts/apply_rules.py (chạy định kỳ ~10 phút) quét log, gắn rule ID (R-001..R-015), "
       "ghi findings vào index siem-findings — với key {rule, src_ip, window} ổn định để khử trùng lặp.")
bullet("triage.py quản lý vòng đời case (open → investigating → escalated → resolved).")
bullet("daily_report.py tự sinh báo cáo Markdown/HTML hàng ngày.")
bullet("Kibana dashboard hiển thị findings thành quyết định giám sát.")

# ---- 3. Tính năng ----
h1("3. Tính năng chính")
for f in [
    "Ingest-pipeline enrichment: tự phân loại attack class tự động cho mọi request.",
    "15 detection rules kiểu Sigma, ánh xạ MITRE ATT&CK.",
    "Tương quan đa nguồn (multi-source): nginx + UFW firewall + SSH auth (R-013/R-014/R-015).",
    "Hàng đợi case bền vững (siem-findings) — hết cảnh báo trùng lặp.",
    "Triage workflow + auto false-positive pass.",
    "Dashboard Kibana + báo cáo daily tự động.",
    "Giám sát liên tục qua cron/Task Scheduler.",
]:
    bullet(f)

# ---- 4. Detection rules ----
h1("4. Detection Rules (R-001…R-015)")
para("15 rule bao phủ từ chữ ký đơn-request đến tương quan tần suất, chuỗi tấn công và đa nguồn. "
     "Tất cả ánh xạ MITRE ATT&CK.")

rules = [
    ("R-001","SQL Injection","3","sqli","nginx"),
    ("R-002","XSS Attempt","2","xss","nginx"),
    ("R-003","Path Traversal / LFI","3","path_traversal","nginx"),
    ("R-004","Command Injection","4","command_injection","nginx"),
    ("R-005","Sensitive File Disclosure","4","info_disclosure","nginx"),
    ("R-006","SSRF","3","ssrf","nginx"),
    ("R-007","Scanner UA","1","scanner","nginx"),
    ("R-008","Admin Panel Probe","2","admin_panel_scan","nginx"),
    ("R-009","Auth Endpoint Scan","2","auth_scan","nginx"),
    ("R-010","Geo-block bypass probe","1","geoblock_bypass","nginx"),
    ("R-011","High Request Rate (freq)","2","transverse_freq","nginx"),
    ("R-012","Attack Chain correlation","3","attack_chain","nginx"),
    ("R-013","Firewall+Web correlation","2-3","cross_source_correlation","UFW + nginx"),
    ("R-014","SSH Brute-Force","3","ssh_bruteforce","SSH auth"),
    ("R-015","SSH Cross-Source correlation","1-3","ssh_cross_source","SSH + web/fw"),
]
tab = doc.add_table(rows=1, cols=5)
tab.style = "Light Grid Accent 1"
tab.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr = tab.rows[0].cells
for i, htxt in enumerate(["ID","Title","Sev","Class","Source"]):
    hdr[i].text = htxt
    for p in hdr[i].paragraphs:
        for rr in p.runs: rr.bold = True
for row in rules:
    cells = tab.add_row().cells
    for i, val in enumerate(row):
        cells[i].text = val
para("")
para("Giá trị gia tăng cốt lõi: tương quan — đơn nguồn (R-011/R-012) và đa nguồn (R-013/R-015) "
     "xuyên nginx + UFW + SSH, cùng hàng đợi triage bền vững mà chỉ WAF lớp-7 không cung cấp được.")

# ---- 5. Dashboard ----
h1("5. Kibana Dashboard (SOC_OVERVIEW)")
para("Dashboard gồm 7 panel kiểu pie/tagcloud (hạn chế render của Kibana 8.12 qua API): "
     "Severity, Rule, Attack Class, MITRE Tactic, Status, SSH Top IPs, UFW Top IPs — cửa sổ 24h.")
para("Số liệu mẫu (24h gần nhất):", bold=True)
bullet("Severity: Critical 7, High 14, Medium 9, Low 1 (tổng 31 findings) — 68% ở mức High/Critical.")
bullet("Rule nổi bật: R-014 SSH brute-force (14), R-005 info disclosure (7), R-008 admin-panel scan (5).")
bullet("MITRE Tactic: credential-access (18) chiếm ưu thế.")
bullet("Status: 31 findings đều open (chưa triage).")
bullet("SSH top IP: 45.148.10.157 (125 lần login failure).")
bullet("UFW top IP: 13.75.39.76 (38 lần bị chặn).")

# ---- 6. Thành phần repo ----
h1("6. Cấu trúc repo")
for line in [
    "AGENTS.md — hướng dẫn cho AI coding agent",
    "ingest-pipelines/nginx-soc-enrich.json — pipeline phân loại tấn công + ECS + GeoIP",
    "rules/detection-rules.yaml — 15 rule Sigma-style",
    "transforms/freq-detection-by-ip.json — R-011 frequency transform",
    "scripts/ — deploy.sh, apply_rules.py, triage.py, daily_report.py, build_dashboard.py, alerts.py (draft), run_soc_detect.bat",
    "dashboards/ — soc_overview.ndjson + SOC_OVERVIEW.md",
    "reports/ — báo cáo hàng ngày và báo cáo tổng thể",
]:
    bullet(line)

# ---- 7. Roadmap / hạn chế ----
h1("7. Roadmap & hạn chế đã biết")
for item in [
    "Backend-direct bypass: log chỉ bắt traffic qua nginx (80/443); tấn công trực tiếp vào backend (DVWA:8080/JuiceShop) không vào nginx-access. Hướng: cài Filebeat trong container backend.",
    "Enrichment backfill: log cũ chưa được enrich ngược; chỉ enrich từ thời điểm chạy.",
    "Alerting: alerts.py là bản nháp — cần nối Slack/Telegram webhook (SOC_ALERT_* env).",
    "Dashboard: chỉ pie/tagcloud qua API; bar/line/metric cần vẽ tay trên UI Kibana.",
]:
    bullet(item)

# ---- 8. Kết luận ----
h1("8. Kết luận")
para("SOC-ify hoàn thành mục tiêu là một SOC thu nhỏ dạy được và có thể tái tạo trên "
     "Azure + Elasticsearch 8.12 + Kibana, không phụ thuộc giấy phép trả phí. Hệ thống đi "
     "từ phát hiện (15 rule + tương quan đa nguồn) đến triage và báo cáo, là nền tảng học "
     "SIEM/SOC và portfolio công việc vững chắc. Không phải SIEM sản xuất thương mại.")

doc.save(OUT)
print("Saved:", os.path.abspath(OUT))
